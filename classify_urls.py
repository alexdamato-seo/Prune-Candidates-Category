#!/usr/bin/env python3
"""
Reads URLs from Column A of an Excel file, scrapes visible text,
classifies each page into one of four categories, and writes
a 'Category' column back to the file.
"""

import re
import time
import openpyxl
import requests
from bs4 import BeautifulSoup

EXCEL_PATH = "/home/user/Prune-Candidates-Category/Prune or Noindex Candidates Landing Folder.xlsx"

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}

# ── keyword lists ────────────────────────────────────────────────────────────

LIFE_SCIENCE_KEYWORDS = [
    "biology", "genomics", "proteomics", "cell biology", "biochemistry",
    "molecular biology", "genetics", "gene", "genome", "protein", "enzyme",
    "dna", "rna", "sequencing", "pcr", "western blot", "elisa", "antibody",
    "cell culture", "microbiome", "bioinformatics", "transcriptomics",
    "metabolomics", "flow cytometry", "microscopy", "assay", "biomarker",
    "stem cell", "crispr", "immunology", "microbiology", "virology",
    "bacteria", "virus", "pathogen", "biopharma", "biotech", "life science",
    "next generation sequencing", "ngs", "mass spectrometry", "chromatography",
]

CLINICAL_KEYWORDS = [
    "clinical", "medicine", "diagnostics", "pharmaceutical", "clinical trial",
    "healthcare", "patient", "hospital", "therapy", "treatment", "disease",
    "drug", "fda", "regulatory", "gmp", "ivd", "in vitro diagnostic",
    "pathology", "oncology", "cardiology", "neurology", "radiology",
    "imaging", "physician", "physician", "medical device", "prognosis",
    "diagnosis", "bioanalytical", "plasma", "serum", "blood", "urine",
    "clinical laboratory", "point of care", "poc", "companion diagnostic",
    "bioequivalence", "pharmacokinetics", "pharmacodynamics",
]

MATERIAL_SCIENCE_KEYWORDS = [
    "polymer", "metal", "composite", "coating", "nanomaterial", "surface analysis",
    "materials science", "material science", "metallurgy", "ceramic",
    "alloy", "thin film", "nanoparticle", "nanotechnology", "tribology",
    "corrosion", "adhesion", "hardness", "tensile", "rheology", "viscosity",
    "xrd", "sem", "tem", "afm", "xps", "esca", "eds", "edx",
    "ftir", "raman", "thermal analysis", "dsc", "tga", "tma", "dma",
    "injection molding", "extrusion", "additive manufacturing", "3d printing",
    "semiconductor", "photovoltaic", "solar cell", "battery", "fuel cell",
    "substrate", "film deposition", "sputtering", "cvd", "pvd",
    "rubber", "plastic", "resin", "epoxy", "silicone", "polyurethane",
    "surface", "interface", "wettability", "contact angle",
]


def scrape_text(url: str, timeout: int = 15) -> tuple[str, str]:
    """Return (visible_text, status) where status is 'ok' or an error message."""
    try:
        resp = requests.get(url, headers=HEADERS, timeout=timeout, allow_redirects=True)
        resp.raise_for_status()
    except requests.exceptions.RequestException as exc:
        return "", f"FAILED: {exc}"

    soup = BeautifulSoup(resp.text, "lxml")

    # Remove non-visible elements
    for tag in soup(["script", "style", "noscript", "head", "meta", "link"]):
        tag.decompose()

    text = soup.get_text(separator=" ", strip=True)
    # Collapse whitespace
    text = re.sub(r"\s+", " ", text).lower()
    return text, "ok"


def count_keyword_hits(text: str, keywords: list[str]) -> int:
    hits = 0
    for kw in keywords:
        hits += len(re.findall(r"\b" + re.escape(kw) + r"\b", text))
    return hits


def classify(text: str) -> str:
    if not text:
        return "Other"

    scores = {
        "Life Science": count_keyword_hits(text, LIFE_SCIENCE_KEYWORDS),
        "Clinical": count_keyword_hits(text, CLINICAL_KEYWORDS),
        "Material Science": count_keyword_hits(text, MATERIAL_SCIENCE_KEYWORDS),
    }

    best_cat = max(scores, key=scores.get)
    best_score = scores[best_cat]

    if best_score == 0:
        return "Other"

    return best_cat


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # Find the header row and URL column
    url_col_idx = None
    header_row = 1

    for cell in ws[header_row]:
        if cell.value and str(cell.value).strip().lower() in ("url", "urls", "link", "links", "page url"):
            url_col_idx = cell.column
            break

    # If no header found, assume column A contains URLs (starting row 1)
    if url_col_idx is None:
        url_col_idx = 1  # Column A

    # Find or create 'Category' column
    max_col = ws.max_column
    cat_col_idx = None

    for cell in ws[header_row]:
        if cell.value and str(cell.value).strip().lower() == "category":
            cat_col_idx = cell.column
            break

    if cat_col_idx is None:
        cat_col_idx = max_col + 1
        ws.cell(row=header_row, column=cat_col_idx, value="Category")

    # Determine data start row
    first_cell_val = ws.cell(row=1, column=url_col_idx).value
    if first_cell_val and str(first_cell_val).strip().lower() in (
        "url", "urls", "link", "links", "page url"
    ):
        data_start_row = 2
    else:
        data_start_row = 1

    failed_urls = []
    processed = 0

    for row_idx in range(data_start_row, ws.max_row + 1):
        url_cell = ws.cell(row=row_idx, column=url_col_idx)
        url = url_cell.value

        if not url or not str(url).strip():
            continue

        url = str(url).strip()
        if not url.startswith(("http://", "https://")):
            url = "https://" + url

        print(f"[{row_idx}] Scraping: {url}", flush=True)

        text, status = scrape_text(url)

        if status != "ok":
            category = "Other"
            failed_urls.append((row_idx, url, status))
            print(f"  -> {status}", flush=True)
        else:
            category = classify(text)
            print(f"  -> {category}", flush=True)

        ws.cell(row=row_idx, column=cat_col_idx, value=category)
        processed += 1

        time.sleep(0.5)  # be polite

    wb.save(EXCEL_PATH)
    print(f"\nDone. Processed {processed} URLs.")

    if failed_urls:
        print(f"\nFailed URLs ({len(failed_urls)}):")
        for row_idx, url, reason in failed_urls:
            print(f"  Row {row_idx}: {url}  ({reason})")
    else:
        print("All URLs loaded successfully.")


if __name__ == "__main__":
    main()
