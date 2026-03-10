#!/usr/bin/env python3
"""
Classifies URLs in Column A of an Excel file into one of four categories
using URL path keywords — no web scraping required.

Categories:
  Life Science    – biology, genomics, proteomics, cell biology, biochemistry
  Clinical        – medicine, diagnostics, pharmaceuticals, clinical, healthcare
  Material Science – polymers, metals, composites, coatings, nanomaterials, surface analysis
  Other           – anything that doesn't clearly fit the above
"""

import re
import openpyxl

EXCEL_PATH = "/home/user/Prune-Candidates-Category/Prune or Noindex Candidates Landing Folder.xlsx"

# ── Rule sets (evaluated in order; first match wins) ────────────────────────
# Each rule is (category, list_of_substring_patterns).
# The path segment checked is everything after the language prefix, lowercased.

RULES = [
    # ── Material Science ────────────────────────────────────────────────────
    ("Material Science", [
        "materialograph",          # materialography and sub-pages
        "pcb-manufactur",          # PCB manufacturing
        "semiconductor-manufactur",# semiconductor manufacturing
        "lithium-ion-battery",     # Li-ion battery
        "microscope-solutions-for-ev",  # electric vehicle manufacturing
        "solid-state-battery",
        "gear-and-bearing",
        "motor-manufactur",
        "power-control-unit",
        "5g-inspection",
        "advancedopticalmetrology",
        "measurement-solutions",
        "microscope-observation-inspection-and-roughness-measurement-solutions-for-medical-device",
        "visit-the-new-evident-industrial-microscopy-labs",
        "trumpf",                  # industrial laser company collab
    ]),

    # ── Clinical ────────────────────────────────────────────────────────────
    ("Clinical", [
        "pathology",
        "medical-device",
        "clinical",
        "diagnostic",
        "pharmaceutical",
        "healthcare",
    ]),

    # ── Life Science ────────────────────────────────────────────────────────
    ("Life Science", [
        "organoid",
        "cell-culture",
        "cell-phenotyping",
        "biomarker",
        "drug-discovery",
        "life-science",
        "bioscapes",               # biological microscopy photography contest
        "ioty",                    # Image of the Year – biological microscopy
        "fv3000",                  # laser confocal fluorescence microscope
        "fv4000",
        "fv5000",
        "ixplore-live-for-luminescence",
        "objectives/research",     # research objectives (cellculture, tirf, livecell, etc.)
        "society-for-neuroscience",
        "neuroscience",
        "truai",                   # AI for fluorescence microscopy
        "truresolution",           # super-resolution fluorescence microscopy
        "trusight",                # spectral imaging
        "truspectral",             # spectral unmixing for fluorescence
        "back-to-science",
        "new-investigator-program",
        "covid-19",
        "return-to-lab",
        "your-science-matters",
    ]),
]

# Anything not matched above → Other


def extract_path(url: str) -> str:
    """Return the lowercased path after stripping the domain and language prefix."""
    url = str(url).strip().lower()
    # Remove protocol + domain
    url = re.sub(r"^https?://[^/]+", "", url)
    # Remove leading language segment like /en/ /de/ /zh/ /ja/ /fr/ /es/ /ko/ /it/
    url = re.sub(r"^/[a-z]{2}/", "/", url)
    return url


def classify_url(url: str) -> str:
    path = extract_path(url)
    for category, patterns in RULES:
        for pat in patterns:
            if pat in path:
                return category
    return "Other"


def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # Detect header row
    header = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
    data_start = 2  # assume row 1 is headers

    # Find URL column (Column A = index 0 in header list, column index 1 in openpyxl)
    url_col = 1  # Column A

    # Find or create 'Category' column
    cat_col = None
    for idx, h in enumerate(header, start=1):
        if h == "category":
            cat_col = idx
            break
    if cat_col is None:
        cat_col = ws.max_column + 1
        ws.cell(row=1, column=cat_col, value="Category")

    results = {"Life Science": 0, "Clinical": 0, "Material Science": 0, "Other": 0}

    for row_idx in range(data_start, ws.max_row + 1):
        url = ws.cell(row=row_idx, column=url_col).value
        if not url or not str(url).strip():
            continue

        category = classify_url(str(url))
        ws.cell(row=row_idx, column=cat_col, value=category)
        results[category] += 1
        print(f"[{row_idx:>3}] {category:<16}  {str(url)[:90]}")

    wb.save(EXCEL_PATH)

    total = sum(results.values())
    print(f"\n{'='*60}")
    print(f"Done — {total} URLs classified")
    print(f"{'='*60}")
    for cat, count in results.items():
        print(f"  {cat:<18} {count:>4}  ({count/total*100:.1f}%)")


if __name__ == "__main__":
    main()
